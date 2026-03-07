"""Test: mapper using key_selector candidates + section title.

For each chunk, LLM receives:
  - section title (scanned from paragraphs above)
  - each placeholder with its context
  - only top 10 candidate keys+values per placeholder (from key_selector)
"""
import sys, io, json, os, re

if not isinstance(sys.stdout, io.TextIOWrapper) or sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from dotenv import load_dotenv
load_dotenv()

import anthropic
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph as P
from key_selector import top_keys


def find_section_title(docx_path: str, para_index: int, max_look_back: int = 20) -> str:
    """Scan paragraphs above para_index for a section title (FORMULAR, uppercase, etc.)."""
    doc = Document(docx_path)
    body = doc.element.body
    paras = [el for el in body if el.tag == qn('w:p')]

    for i in range(para_index - 1, max(0, para_index - max_look_back) - 1, -1):
        para = P(paras[i], None)
        txt = para.text.strip()
        if not txt or len(txt) < 3:
            continue
        if 'formular' in txt.lower():
            # Grab this line + next non-empty line as title
            title = txt
            for j in range(i + 1, min(i + 4, len(paras))):
                next_txt = P(paras[j], None).text.strip()
                if next_txt and len(next_txt) > 3:
                    title += ' - ' + next_txt
                    break
            return title
        upper_ratio = sum(1 for c in txt if c.isupper()) / max(len(txt), 1)
        if upper_ratio > 0.6 and len(txt) > 5 and len(txt) < 100:
            return txt
    return ''


def build_page_text_with_candidates(docx_path: str, p_start: int, p_end: int,
                                     fields: list, data: dict, n_cand: int = 10):
    """Build page text with [[PH:id]] markers + per-placeholder candidate lists."""
    doc = Document(docx_path)
    body = doc.element.body
    paras = [el for el in body if el.tag == qn('w:p')]

    fields_by_para = {}
    for fld in fields:
        if fld.get('field_type') == 'CHECKBOX':
            continue
        loc = fld.get('location', '')
        m = re.findall(r'p:(\d+)', loc)
        if m:
            pidx = int(m[-1])
            if p_start <= pidx <= p_end:
                fields_by_para.setdefault(pidx, []).append(fld)

    page_fields = []
    lines = []

    for i in range(p_start, min(p_end + 1, len(paras))):
        para = P(paras[i], None)
        txt = para.text.strip()
        if not txt:
            continue
        if i in fields_by_para:
            flds = sorted(fields_by_para[i], key=lambda f: f['start'])
            marked = txt
            offset = 0
            for fld in flds:
                s = fld['start'] + offset
                e = fld['end'] + offset
                fid = fld['field_id'][:8]
                marker = f'[[PH:{fid}]]'
                marked = marked[:s] + marker + marked[e:]
                offset += len(marker) - (fld['end'] - fld['start'])
                page_fields.append(fld)
            lines.append(marked)
        else:
            lines.append(txt)

    page_text = '\n'.join(lines)

    # Build candidate lists per field
    cand_section = "\nCANDIDATE KEYS PER PLACEHOLDER (pick from these):\n"
    for fld in page_fields:
        fid = fld['field_id'][:8]
        candidates = top_keys(fld, data, n=n_cand)
        cand_lines = ', '.join(f'"{k}": "{v}"' for k, _, v in candidates)
        cand_section += f"  {fid}: {{{cand_lines}}}\n"

    return page_text, page_fields, cand_section


def map_page_with_ks(page_text: str, page_fields: list, cand_section: str,
                     section_title: str, model: str = 'claude-haiku-4-5-20251001') -> dict:
    """Send page text + candidate keys to LLM."""

    prompt = f"""You are a document data-entry system for Romanian procurement forms.

SECTION TITLE: {section_title}

Below is a page from the form with placeholders marked as [[PH:xxxxxxxx]].
For each placeholder, select the BEST matching key from its candidate list below.

RULES:
1. Select an exact key from the placeholder's candidate list, or null if none fits.
2. Use the FULL page context to understand which key belongs where.
3. When the same structure repeats (e.g. first = "Asociat 1", second = "Asociat 2"), use ordering.
4. If a placeholder asks for PART of a value (e.g. just name from "name, role"), provide that part in extracted_value.
5. Return ONLY a JSON array.

DOCUMENT PAGE:
{page_text}
{cand_section}
Return ONLY a JSON array:
[{{"id": "xxxxxxxx", "selected_key": "key"|null, "extracted_value": "substring"|null, "confidence": 0.0-1.0, "reasoning": "brief"}}]
"""

    client = anthropic.Anthropic(api_key=os.environ['ANTHROPIC_API_KEY'])
    resp = client.messages.create(
        model=model, max_tokens=8192, temperature=0,
        messages=[{"role": "user", "content": prompt}],
    )
    content = (resp.content[0].text or "").strip()
    content = re.sub(r"^```(?:json)?\s*", "", content)
    content = re.sub(r"\s*```$", "", content).strip()
    m = re.search(r"(\[[\s\S]*\])", content)
    items = json.loads(m.group(1)) if m else json.loads(content)

    result = {}
    for item in items:
        eid = item.get("id")
        if eid:
            result[eid] = item
    return result


if __name__ == '__main__':
    with open('parser_result.json', encoding='utf-8') as f:
        pr = json.load(f)
    with open('input_date.json', encoding='utf-8') as f:
        data = json.load(f)

    # FORMULAR 8: p:350 (FORMULAR 8 title) to p:417
    title = find_section_title('sample_forms.docx', 359)
    print(f"Section title: {title}")

    page_text, page_fields, cand_section = build_page_text_with_candidates(
        'sample_forms.docx', 350, 417, pr['fields'], data
    )
    print(f"Fields: {len(page_fields)}")
    print(f"Page text: {len(page_text)} chars\n")

    print("Calling LLM...\n")
    result = map_page_with_ks(page_text, page_fields, cand_section, title)

    print(f"Mapped {len(result)} fields:\n")
    for fld in page_fields:
        fid8 = fld['field_id'][:8]
        loc = fld.get('location', '')
        label = fld.get('label', '')[:40]
        r = result.get(fid8, {})
        key = r.get('selected_key', '???')
        ext = r.get('extracted_value')
        conf = r.get('confidence', 0)
        ext_str = f' -> "{ext}"' if ext else ''
        print(f"  {loc:15s} [{label:40s}] => {key}{ext_str}  ({conf})")
